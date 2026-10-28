"""Build the printable Bible Outlines booklet from the Markdown in docs/.

The original Word booklet (booklet/src/template.docx) supplies the cover page,
styles, bullet numbering, footer, and page size. Everything after the cover is
regenerated from the pages listed in mkdocs.yml's nav, so the website and the
booklet always share one source.

Booklet-only layout hints are HTML comments, which the website ignores:

    <!-- pagebreak -->            start a new page
    <!-- colbreak -->             jump to the next column
    <!-- cols:N -->               new section with N columns, on a new page
    <!-- cols:N continuous -->    new section with N columns, same page
    <!-- docx-image: PATH width=2.8in -->   picture (PATH relative to the page)
    <!-- docx-skip --> ... <!-- /docx-skip -->   website-only content

Usage:  python scripts/build_booklet.py [--out "booklet/Bible Outlines.docx"]
"""

import argparse
import copy
import math
import re
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
TEMPLATE = ROOT / "booklet" / "src" / "template.docx"
DEFAULT_OUT = ROOT / "booklet" / "Bible Outlines.docx"

COVER_TITLE = "Old & New Testament"  # replaces "Old Testament" on the cover
BULLET_NUM_ID = "26"  # the template's bullet list (List Paragraph style)
PAGE_TEXT_WIDTH = 8640  # twips: 7in page minus 0.5in margins
COL_SPACE = 360
ACCENT = "3F51B5"

DIRECTIVE = re.compile(r"^<!--\s*(.*?)\s*-->$")
LIST_ITEM = re.compile(r"^( *)([-*]|\d+\.) (.*)$")
ADMONITION = re.compile(r'^(!!!|\?\?\?\+?) *(\w+)(?: +"(.*)")?\s*$')


# --------------------------------------------------------------------------
# Markdown (the subset used in docs/) -> blocks
# --------------------------------------------------------------------------

def parse_blocks(lines):
    """Turn Markdown lines into a flat list of (kind, data) blocks."""
    blocks, i = [], 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        m = DIRECTIVE.match(stripped)
        if m:
            if m.group(1) == "docx-skip":
                while i < len(lines) and lines[i].strip() != "<!-- /docx-skip -->":
                    i += 1
            else:
                blocks.append(("directive", m.group(1)))
            i += 1
            continue
        if stripped.startswith("<div") or stripped.startswith("</div"):
            # Website-only card grids (home page); skip to the closing tag.
            while i < len(lines) and not lines[i].strip().startswith("</div"):
                i += 1
            i += 1
            continue
        m = re.match(r"^(#{1,4}) (.*)$", line)
        if m:
            blocks.append(("h%d" % len(m.group(1)), m.group(2).strip()))
            i += 1
            continue
        m = ADMONITION.match(stripped)
        if m:
            body = []
            i += 1
            while i < len(lines) and (lines[i].startswith("    ") or not lines[i].strip()):
                body.append(lines[i][4:])
                i += 1
            title = m.group(3) if m.group(3) is not None else m.group(2).capitalize()
            blocks.append(("admonition", (title, parse_blocks(body))))
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        if LIST_ITEM.match(line):
            items = []
            while i < len(lines) and LIST_ITEM.match(lines[i]):
                indent, marker, text = LIST_ITEM.match(lines[i]).groups()
                items.append((len(indent) // 4, marker, text))
                i += 1
            blocks.append(("list", items))
            continue
        if stripped.startswith(">"):
            text = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                text.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("quote", " ".join(text)))
            continue
        text = []
        while i < len(lines) and lines[i].strip() and not _starts_block(lines[i]):
            text.append(lines[i].strip())
            i += 1
        blocks.append(("para", " ".join(text)))
    return blocks


def _starts_block(line):
    s = line.strip()
    return bool(
        DIRECTIVE.match(s) or re.match(r"^#{1,4} ", line) or ADMONITION.match(s)
        or s.startswith("|") or LIST_ITEM.match(line) or s.startswith(">")
    )


def read_page(path):
    lines = path.read_text(encoding="utf-8").splitlines()
    if lines and lines[0].strip() == "---":  # YAML front matter
        end = lines.index("---", 1)
        lines = lines[end + 1:]
    return parse_blocks(lines)


def clean_inline(text):
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)  # links -> text
    text = re.sub(r":[a-z][a-z0-9]*(?:-[a-z0-9]+)+:", "", text)  # emoji shortcodes
    text = re.sub(r"\+\+(\w+)\+\+", r"\1", text)  # keyboard keys
    return text


def inline_runs(text):
    """Yield (text, bold, italic, line_break) from **bold** / *italic* / <br>."""
    bold = italic = False
    for tok in re.split(r"(\*\*\*|\*\*|\*|<br>)", clean_inline(text)):
        if tok == "***":
            bold, italic = not bold, not italic
        elif tok == "**":
            bold = not bold
        elif tok == "*":
            italic = not italic
        elif tok == "<br>":
            yield ("", False, False, True)
        elif tok:
            yield (tok, bold, italic, False)


def plain(text):
    return "".join(t for t, *_ in inline_runs(text))


# --------------------------------------------------------------------------
# Word output
# --------------------------------------------------------------------------

PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
    "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
    "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
    "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
    "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
    "sectPr", "pPrChange",
]


def set_ppr_child(paragraph, element):
    """Insert a pPr child in schema order, replacing any existing one."""
    pPr = paragraph._p.get_or_add_pPr()
    tag = element.tag.split("}")[1]
    for old in pPr.findall(qn("w:" + tag)):
        pPr.remove(old)
    rank = PPR_ORDER.index(tag)
    for idx, child in enumerate(pPr):
        child_tag = child.tag.split("}")[1]
        if child_tag in PPR_ORDER and PPR_ORDER.index(child_tag) > rank:
            pPr.insert(idx, element)
            return
    pPr.append(element)


def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


class Booklet:
    def __init__(self, template):
        self.doc = Document(str(template))
        body = self.doc.element.body
        children = list(body.iterchildren())
        cover_end = next(
            i for i, c in enumerate(children)
            if c.tag == qn("w:p") and c.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None
        )
        self.final_sectpr = children[-1]
        for c in children[cover_end + 1:-1]:
            body.remove(c)
        for p in self.doc.paragraphs[:cover_end]:
            for t in p._p.iter(qn("w:t")):
                if t.text == "Old Testament":
                    t.text = COVER_TITLE
        self.section = None  # {"cols": n, "type": "nextPage"|"continuous"}
        self.at_top = True
        self.break_before = False
        self.border = False

    # -- sections ---------------------------------------------------------

    def _sectpr(self, cols, kind):
        s = copy.deepcopy(self.final_sectpr)
        for t in s.findall(qn("w:type")):
            s.remove(t)
        if kind == "continuous":
            s.insert(0, el("w:type", val="continuous"))
        c = s.find(qn("w:cols"))
        for a in list(c.attrib):
            del c.attrib[a]
        c.set(qn("w:space"), str(COL_SPACE))
        if cols > 1:
            c.set(qn("w:num"), str(cols))
        return s

    def start_section(self, cols, kind="nextPage"):
        if self.section is not None:
            if kind == "continuous" and self.section["cols"] == cols:
                return
            p = self.doc.add_paragraph()
            set_ppr_child(p, self._sectpr(self.section["cols"], self.section["type"]))
        self.section = {"cols": cols, "type": kind}
        if kind == "nextPage":
            # A new-page section already starts a page; a pending break would add a blank one.
            self.at_top = True
            self.break_before = False

    def finish(self):
        s = self._sectpr(self.section["cols"], self.section["type"])
        body = self.doc.element.body
        body.replace(self.final_sectpr, s)

    @property
    def width(self):
        cols = self.section["cols"]
        return (PAGE_TEXT_WIDTH - COL_SPACE * (cols - 1)) // cols

    # -- paragraphs ---------------------------------------------------------

    def para(self, text="", style=None, size=None, bold=None, italic=None):
        p = self.doc.add_paragraph(style=style)
        if self.break_before:
            p.paragraph_format.page_break_before = True
            self.break_before = False
        if text:
            self.add_runs(p, text, size=size, bold=bold, italic=italic)
        if self.border:
            bdr = el("w:pBdr")
            bdr.append(el("w:left", val="single", sz=18, space=6, color=ACCENT))
            set_ppr_child(p, bdr)
            p.paragraph_format.left_indent = p.paragraph_format.left_indent or Pt(6)
        self.at_top = False
        return p

    @staticmethod
    def add_runs(p, text, size=None, bold=None, italic=None):
        for t, b, i, br in inline_runs(text):
            r = p.add_run()
            if br:
                r.add_break()
                continue
            r.text = t
            r.bold = True if bold else (b or None)
            r.italic = True if italic else (i or None)
            if size:
                r.font.size = Pt(size)

    def page_break(self):
        # Applied as "page break before" on the next paragraph: a separate break
        # paragraph can itself spill onto a fresh page and leave it blank.
        if not self.at_top:
            self.break_before = True
            self.at_top = True

    def column_break(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.COLUMN)

    def image(self, page_path, spec):
        m = re.match(r"(\S+)(?:\s+width=([\d.]+)in)?", spec)
        path = (page_path.parent / m.group(1)).resolve()
        width = min(float(m.group(2) or 2.8), self.width / 1440)
        p = self.para()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))

    def table(self, rows):
        if self.break_before:
            self.para()
        ncols = max(len(r) for r in rows)
        rows = [r + [""] * (ncols - len(r)) for r in rows]
        weights = [
            max(math.sqrt(max(len(plain(r[c])) for r in rows)), 2.0) for c in range(ncols)
        ]
        total = self.width
        widths = [int(total * w / sum(weights)) for w in weights]
        t = self.doc.add_table(rows=len(rows), cols=ncols)
        t.style = "Table Grid"
        t.autofit = False
        tblPr = t._tbl.tblPr
        tblPr.append(el("w:tblW", w=total, type="dxa"))
        cell_mar = el("w:tblCellMar")
        for side in ("left", "right"):
            cell_mar.append(el("w:" + side, w=60, type="dxa"))
        tblPr.append(cell_mar)
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))
        for ri, row in enumerate(rows):
            for ci, text in enumerate(row):
                cell = t.cell(ri, ci)
                tcPr = cell._tc.get_or_add_tcPr()
                tcPr.append(el("w:tcW", w=widths[ci], type="dxa"))
                if ri == 0:
                    tcPr.append(el("w:shd", val="clear", color="auto", fill="E8EAF6"))
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                self.add_runs(p, text, size=8.5, bold=(ri == 0) or None)
        if len(rows) > 1:
            tr_pr = t.rows[0]._tr.get_or_add_trPr()
            tr_pr.append(el("w:tblHeader"))
        # Keep a little air after each table.
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        self.at_top = False

    def list_items(self, items):
        for level, marker, text in items:
            if marker[0].isdigit():
                p = self.para(style="Normal")
                p.paragraph_format.left_indent = Inches(0.3 + 0.25 * level)
                p.paragraph_format.first_line_indent = Inches(-0.2)
                self.add_runs(p, f"{marker} {text}")
            else:
                p = self.para(text, style="List Paragraph")
                if level:
                    numPr = el("w:numPr")
                    numPr.append(el("w:ilvl", val=min(level, 2)))
                    numPr.append(el("w:numId", val=BULLET_NUM_ID))
                    set_ppr_child(p, numPr)

    # -- blocks ---------------------------------------------------------------

    def render(self, blocks, page_path):
        after_h2 = after_h1 = False
        for kind, data in blocks:
            if after_h1 and not (kind == "directive" and data.startswith("cols:")):
                self.start_section(2, "continuous")
            after_h1 = False
            if kind == "directive":
                self.directive(data, page_path)
            elif kind == "h1":
                p = self.para(data, style="Title", size=24)
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.keep_with_next = True
                after_h1 = True
            elif kind == "h2":
                p = self.para(style="Heading 1")
                self.add_runs(p, data, size=14)
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(2)
            elif kind == "para" and after_h2 and re.fullmatch(r"\*[^*].*[^*]\*", data):
                self.para(data[1:-1], style="Subtitle")
            elif kind in ("h3", "h4"):
                p = self.para(data, bold=True, italic=(kind == "h4") or None)
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.keep_with_next = True
            elif kind == "para":
                p = self.para(data)
                p.paragraph_format.space_after = Pt(4)
            elif kind == "quote":
                p = self.para(data)
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after = Pt(4)
            elif kind == "list":
                self.list_items(data)
            elif kind == "table":
                self.table(data)
            elif kind == "admonition":
                title, inner = data
                self.border = True
                p = self.para(title, bold=True, italic=True)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.keep_with_next = True
                self.render(inner, page_path)
                self.border = False
                gap = self.doc.add_paragraph()
                gap.paragraph_format.space_after = Pt(2)
            after_h2 = kind == "h2"

    def directive(self, text, page_path):
        if text == "pagebreak":
            self.page_break()
        elif text == "colbreak":
            self.column_break()
        elif text.startswith("cols:"):
            parts = text[5:].split()
            self.start_section(int(parts[0]), parts[1] if len(parts) > 1 else "nextPage")
        elif text.startswith("docx-image:"):
            self.image(page_path, text[len("docx-image:"):].strip())

    def add_page(self, path):
        self.start_section(1, "nextPage")
        self.render(read_page(path), path)


def nav_pages():
    """Page paths in mkdocs.yml nav order, skipping the home page."""
    loader = yaml.SafeLoader
    loader.add_multi_constructor("tag:yaml.org,2002:python/", lambda l, s, n: None)
    config = yaml.load((ROOT / "mkdocs.yml").read_text(encoding="utf-8"), Loader=loader)

    def walk(node):
        if isinstance(node, str):
            yield node
        elif isinstance(node, list):
            for item in node:
                yield from walk(item)
        elif isinstance(node, dict):
            for value in node.values():
                yield from walk(value)

    return [DOCS / p for p in walk(config["nav"]) if p != "index.md"]


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()

    booklet = Booklet(TEMPLATE)
    for page in nav_pages():
        booklet.add_page(page)
    booklet.finish()
    args.out.parent.mkdir(parents=True, exist_ok=True)
    booklet.doc.save(str(args.out))
    print(f"Wrote {args.out.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
